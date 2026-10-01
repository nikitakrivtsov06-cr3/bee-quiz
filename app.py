import streamlit as st
import os
import io
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill
import gspread
from google.oauth2.service_account import Credentials

st.set_page_config(page_title="Sunday Quiz", layout="wide")

# --- ПОДКЛЮЧЕНИЕ К GOOGLE SHEETS ---
@st.cache_resource
def get_gsheets_client():
    """Авторизуется через сервисный аккаунт и возвращает клиент gspread."""
    creds_dict = dict(st.secrets["connections"]["gsheets"]["credentials"])
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    client = gspread.authorize(creds)
    return client

def get_worksheet():
    """Возвращает первый лист Google Таблицы."""
    client = get_gsheets_client()
    spreadsheet_url = st.secrets["connections"]["gsheets"]["spreadsheet"]
    sheet = client.open_by_url(spreadsheet_url).sheet1
    return sheet

def save_to_gsheets(answers, comments):
    """Сохраняет ответы пользователя в Google Sheets."""
    try:
        sheet = get_worksheet()
        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            answers.get("q_1", ""),
            comments.get("comment_1", ""),
            answers.get("q_2", ""),
            comments.get("comment_2", ""),
            answers.get("q_3", ""),
            comments.get("comment_3", ""),
        ]
        sheet.append_row(row, value_input_option="USER_ENTERED")
        return True
    except Exception as e:
        st.error(f"Ошибка сохранения в Google Sheets: {e}")
        return False

def read_from_gsheets():
    """Читает все данные из Google Sheets в DataFrame."""
    try:
        sheet = get_worksheet()
        data = sheet.get_all_records()
        return pd.DataFrame(data)
    except Exception as e:
        st.error(f"Ошибка чтения из Google Sheets: {e}")
        return pd.DataFrame()

# --- CSS ---
st.markdown("""
    <style>
    .stApp { 
        background-color: #f4f8fb !important; 
    }
    h1, h2, h3, h4, h5, h6,
    .stMarkdown, .stMarkdown p, .stCaption {
        color: #1a1a1a !important;
    }
    
    div[data-testid="stExpander"] {
        background-color: #ffffff !important;
        border: 1px solid #1a1a1a !important;
        border-radius: 8px !important;
        overflow: hidden !important;
        box-shadow: none !important;
    }
    div[data-testid="stExpander"] summary {
        background-color: #ffffff !important;
        padding: 14px 16px !important;
        cursor: pointer !important;
        list-style: none !important;
    }
    div[data-testid="stExpander"] summary:hover {
        background-color: #f5f5f5 !important;
    }
    div[data-testid="stExpander"] summary * {
        color: #1a1a1a !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
    }
    
    div[role="radiogroup"] {
        gap: 10px !important;
    }
    div[role="radiogroup"] > label {
        background-color: #ffffff !important;
        padding: 14px 18px !important;
        border-radius: 12px !important;
        border: 1px solid #e1e8ed !important;
        box-shadow: 0px 2px 6px rgba(0, 0, 0, 0.03) !important;
        width: 100% !important;
        cursor: pointer !important;
        transition: all 0.2s ease !important;
        display: flex !important;
        align-items: center !important;
    }
    div[role="radiogroup"] > label:hover {
        border-color: #ff4b4b !important;
        background-color: #fff8f8 !important;
    }
    div[role="radiogroup"] > label,
    div[role="radiogroup"] > label *,
    div[role="radiogroup"] > label p,
    div[role="radiogroup"] > label span {
        color: #1a1a1a !important;
        -webkit-text-fill-color: #1a1a1a !important;
        font-size: 1rem !important;
        cursor: pointer !important;
    }
    
    div[role="radiogroup"] > label > div:first-child > div:first-child {
        background-color: #ffffff !important;
        border: 2px solid #b0b8c1 !important;
        border-radius: 50% !important;
    }
    
    .stTextArea textarea {
        color: #1a1a1a !important;
        -webkit-text-fill-color: #1a1a1a !important;
        background-color: #ffffff !important;
        font-size: 1rem !important;
    }
    .stTextArea textarea::placeholder {
        color: #999999 !important;
        -webkit-text-fill-color: #999999 !important;
    }
    
    @media (max-width: 768px) {
        .block-container {
            padding-left: 1rem !important;
            padding-right: 1rem !important;
            padding-top: 1rem !important;
        }
        h2 { 
            font-size: 1.35rem !important; 
            line-height: 1.3 !important;
        }
    }
    </style>
""", unsafe_allow_html=True)

# --- ИНИЦИАЛИЗАЦИЯ ---
if 'step' not in st.session_state:
    st.session_state.step = 1
if 'answers' not in st.session_state:
    st.session_state.answers = {}
if 'comments' not in st.session_state:
    st.session_state.comments = {}

# --- ВОПРОСЫ ---
questions = [
    {
        "title": "Как вам концепция игры про автоматизацию улья?",
        "description": (
            "**Строительство фабрики внутри улья, доведённое до абсолюта.**\n\n"
            "Вы — пчела, которой наскучила ежедневная рутина, поэтому вы решаете "
            "автоматизировать добычу и обработку по максимуму.\n\n"
            "Вашей первичной целью станет постройка небольшого завода по производству мёда "
            "и его транспортировка, причём чем выше эффективность вашего завода, тем больше "
            "последователей (работников) примкнёт к вашему «инженерному улью».\n\n"
            "Вы столкнётесь со сложностями в виде Королевы улья. Ей не нравится, что пчёлы "
            "выходят из её подчинения, выбирая продвинутые методы развития. По мере роста "
            "популяции вы столкнётесь с репрессиями, санкциями и законами, усложняющими "
            "постройку новых механизмов или запрещающими работу на старых.\n\n"
            "Придётся выстраивать логистические цепочки с учётом «завтра», принимать "
            "антикризисные решения и в конечном счёте добиться подавляющего большинства "
            "сторонников в улье!"
        ),
        "options": [
            "Очень нравится",
            "Скорее нравится, чем нет",
            "Ну спорно",
            "Скорее нет, чем да",
            "Вообще нет"
        ],
        "image": "bees.jpg"
    },
    {
        "title": "Как вам концепция квеста в лабиринте на два игрока?",
        "description": (
            "**Квест в лабиринте на два игрока, где один ищет выход, а второй выстраивает его маршрут по камерам в реальном времени... Конечно если один из них не является предателем.**\n\n"
            "Вы оказываетесь в ловушке холодного, погружённого во мрак комплекса. Игра предлагает примерить одну из двух ролей: бегуна, ищущего выход среди десятков слабо освещённых комнат, или оператора, который наблюдает за лабиринтом через сеть уцелевших видеокамер.\n\n"
            "Вашей первичной целью станет совместный побег. Оператору предстоит выстраивать безопасный маршрут в реальном времени, а бегуну — полагаться на его подсказки, чтобы миновать запутанные коридоры, скрытые во тьме тупики и смертоносные механизмы.\n\n"
            "Однако главная опасность кроется не в ловушках, а в голосе на другом конце рации. Вы никогда не можете быть до конца уверены в своём напарнике, ведь любой из вас (а иногда и оба!) может оказаться предателем. Для предателя приоритеты кардинально меняются: вместо спасения команды вам предстоит хладнокровно устранить своего спутника, заманив его в западню или навсегда заперев во тьме, а если вы вдруг узнали что предателем является ваш напарник то вам необходимо как можно быстрее в одиночку выбраться из лабиринта, оставив его там навсегда.\n\n"
            "Вам придётся шаг за шагом преодолевать страх неизвестности, анализировать каждое действие напарника и в конечном счёте решить: готовы ли вы доверить свою жизнь этому человеку чтобы сбежать вместе, или постараетесь сбежать без него и оставите здесь, возможно, навсегда?"
        ),
        "options": [
            "Очень нравится",
            "Скорее нравится, чем нет",
            "Ну спорно",
            "Скорее нет, чем да",
            "Вообще нет"
        ],
        "image": "maze.jpg"
    },
    {
        "title": "Как вам концепция экшен-рогалика с вселением в тела врагов?",
        "description": (
            "**Представьте экшен-рогалик, где у вас нет собственного оружия, а единственный способ двигаться вперед - вселиться в тело врага и управлять им... пока оно не разлетится на куски.**\n\n"
            "Вы играете за бестелесную душу в глубинах древнего подземелья. Могущественный волшебник пал от предательства соратников, лишившись тела и сил. Все что осталось, это хрупкий дух с бессмертным желанием выбраться и отомстить. Когда твоя личность грозит развеяться от любого сквозняка, морали уже нет. И некромантия уже не зло, а единственный способ существования.\n\n"
            "А еще это способ сражаться с врагами. Побеждая монстра, вы можете занять его тело, перенимая его атаки и уникальные способности — будь то тяжелый щит, дальнобойные стрелы или заклинание полета. Конечно, захваченные тела не бессмертны, однако даже получение урона может сыграть вам на руку. Когда у подчиненной тушки кончается здоровье, вы эффектно сбрасываете ее прямо в центре отряда противников, активируя уникальный посмертный эффект. Взрыв, отравляющее облако или область безумия - вариантов много, вы сможете наблюдать их со стороны, перелетев в следующее подходящее тело.\n\n"
            "Геймплей поощряет высокий темп и точный расчет. Выживание зависит от умения на ходу адаптироваться к разным боевым стилям, грамотно выстраивать цепочки целей, а главное: превращать силу врагов в свою собственную. Сможете ли вы освоить мувсет каждого монстра подземелья, собрать идеальный билд и дойти до конца, чтобы сразиться с предателями и открыть тайну собственной смерти?"
        ),
        "options": [
            "Очень нравится",
            "Скорее нравится, чем нет",
            "Ну спорно",
            "Скорее нет, чем да",
            "Вообще нет"
        ],
        "image": "soul.jpg"
    }
]

total_steps = len(questions)
current_step = st.session_state.step

# --- ФУНКЦИЯ ДЛЯ EXCEL ---
def to_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Ответы')
        worksheet = writer.sheets['Ответы']
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='FF4B4B', end_color='FF4B4B', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        worksheet.row_dimensions[1].height = 30
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        cell_len = max(len(str(line)) for line in str(cell.value).split('\n'))
                        max_length = max(max_length, cell_len)
                except:
                    pass
            adjusted_width = min(max_length + 3, 55)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        worksheet.freeze_panes = 'A2'
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    return output.getvalue()

# --- ОСНОВНОЙ ИНТЕРФЕЙС ---
if current_step <= total_steps:
    q_data = questions[current_step - 1]
    
    st.progress((current_step - 1) / total_steps)
    st.caption(f"шаг {current_step}/{total_steps}")
    st.markdown(f"## **{q_data['title']}**")
    st.caption("Выберите один ответ")
    
    if q_data["description"]:
        with st.expander("📖  Нажмите, чтобы прочитать описание игры", expanded=False):
            st.markdown(q_data["description"])
    
    previous_choice = st.session_state.answers.get(f"q_{current_step}")
    if previous_choice in q_data["options"]:
        default_index = q_data["options"].index(previous_choice)
    else:
        default_index = None

    selected_option = st.radio(
        label="Options",
        options=q_data["options"],
        index=default_index,
        key=f"radio_{current_step}",
        label_visibility="collapsed"
    )
    
    comment = st.text_area(
        "Комментарий (необязательно):",
        value=st.session_state.comments.get(f"comment_{current_step}", ""),
        key=f"comment_input_{current_step}",
        placeholder="Здесь можно написать свой комментарий...",
        height=100
    )
    
    image_path = q_data["image"]
    if os.path.exists(image_path):
        st.image(image_path, use_container_width=True)
    else:
        st.image("https://cdn.pixabay.com/photo/2017/01/06/19/15/soap-1958683_1280.jpg", use_container_width=True)
    
    col_prev, col_next = st.columns([1, 1])
    with col_prev:
        if st.button("← Назад", disabled=(current_step == 1), use_container_width=True):
            st.session_state.step -= 1
            st.rerun()
    with col_next:
        if st.button("Далее →", type="primary", disabled=(selected_option is None), use_container_width=True):
            st.session_state.answers[f"q_{current_step}"] = selected_option
            st.session_state.comments[f"comment_{current_step}"] = comment
            
            if current_step == total_steps:
                if save_to_gsheets(st.session_state.answers, st.session_state.comments):
                    st.success("Ответы сохранены!")
            
            st.session_state.step += 1
            st.rerun()

else:
    st.balloons()
    st.write("")
    st.write("")
    
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.success("🎉 Спасибо за прохождение опроса!")
        st.write("")
        if st.button("Пройти заново", type="primary", use_container_width=True):
            st.session_state.step = 1
            st.session_state.answers = {}
            st.session_state.comments = {}
            st.rerun()
    
    st.write("")
    st.write("")

# --- АДМИН-ПАНЕЛЬ ---
query_params = st.query_params
is_admin = query_params.get("admin") == "true"

if is_admin:
    with st.sidebar:
        st.title("Панель администратора")
        admin_pass = st.text_input("Пароль:", type="password")
        
        correct_password = st.secrets.get("admin", {}).get("password", "1234")
        
        if admin_pass == correct_password:
            st.success("Доступ разрешён")
            st.subheader("Сохранённые ответы:")
            
            df = read_from_gsheets()
            if df is not None and not df.empty:
                st.dataframe(df)
                st.caption(f"Всего ответов: **{len(df)}**")
                
                st.write("**Скачать результаты:**")
                date_str = datetime.now().strftime("%Y-%m-%d_%H-%M")
                st.download_button(
                    "📊 Excel",
                    data=to_excel(df),
                    file_name=f"results_{date_str}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                st.download_button(
                    "📄 CSV",
                    data=df.to_csv(index=False, encoding='utf-8-sig'),
                    file_name=f"results_{date_str}.csv",
                    mime="text/csv",
                    use_container_width=True
                )
            else:
                st.info("Ответов пока нет.")
        elif admin_pass:
            st.error("Неверный пароль")
